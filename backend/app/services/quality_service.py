import csv
import io
import logging
from datetime import datetime, timedelta
from typing import Optional

from sqlalchemy import select, func, case, and_, extract
from sqlalchemy.ext.asyncio import AsyncSession

from backend.app.models.models import QualityTraceIndex, Ticket
from backend.app.schemas.quality import QualityTraceQuery, QualityDashboardQuery, QualityExportQuery

logger = logging.getLogger(__name__)


class QualityService:
    """质量追溯服务"""

    async def get_trace_data(self, db: AsyncSession, query: QualityTraceQuery) -> dict:
        """质量追溯查询（按维度聚合）"""
        conditions = []
        if query.model_number:
            conditions.append(QualityTraceIndex.model_number == query.model_number)
        if query.batch_code:
            conditions.append(QualityTraceIndex.batch_code == query.batch_code)
        if query.issue_category:
            conditions.append(QualityTraceIndex.issue_category == query.issue_category)
        if query.action_type:
            conditions.append(QualityTraceIndex.order_type == query.action_type)
        if query.archive_complete is not None:
            conditions.append(QualityTraceIndex.archive_complete == query.archive_complete)
        if query.start_date:
            start_dt = datetime.strptime(query.start_date, "%Y-%m-%d")
            conditions.append(QualityTraceIndex.created_at >= start_dt)
        if query.end_date:
            end_dt = datetime.strptime(query.end_date, "%Y-%m-%d") + timedelta(days=1)
            conditions.append(QualityTraceIndex.created_at < end_dt)

        # 查询汇总数据
        summary_stmt = select(
            func.count().label("total_records"),
            func.count(QualityTraceIndex.ticket_id.distinct()).label("total_tickets"),
            func.count(QualityTraceIndex.ticket_id).label("total_archived"),
            func.avg(QualityTraceIndex.resolution_days).label("avg_resolution_days"),
        )
        if conditions:
            summary_stmt = summary_stmt.where(*conditions)
        summary_result = await db.execute(summary_stmt)
        summary_row = summary_result.one()

        summary = {
            "total_records": summary_row.total_records or 0,
            "total_tickets": summary_row.total_tickets or 0,
            "total_archived": summary_row.total_archived or 0,
            "avg_resolution_days": round(summary_row.avg_resolution_days, 1) if summary_row.avg_resolution_days else 0,
        }

        # 按维度分组聚合
        group_by = query.group_by
        if group_by == "model":
            dimension_col = QualityTraceIndex.model_number
        elif group_by == "batch":
            dimension_col = QualityTraceIndex.batch_code
        elif group_by == "date":
            dimension_col = func.date_format(QualityTraceIndex.created_at, "%Y-%m-%d")
        elif group_by == "category":
            dimension_col = QualityTraceIndex.issue_category
        else:
            dimension_col = QualityTraceIndex.model_number

        # 分组查询
        group_stmt = select(
            dimension_col.label("dimension"),
            func.count().label("record_count"),
            func.count(QualityTraceIndex.ticket_id.distinct()).label("ticket_count"),
            func.count(QualityTraceIndex.ticket_id).label("archived_count"),
            func.avg(QualityTraceIndex.resolution_days).label("avg_resolution_days"),
        ).group_by(dimension_col).order_by(func.count().desc())

        if conditions:
            group_stmt = group_stmt.where(*conditions)

        # 分页
        offset = (query.page - 1) * query.page_size
        group_stmt = group_stmt.offset(offset).limit(query.page_size)

        group_result = await db.execute(group_stmt)
        group_rows = group_result.all()

        groups = []
        for row in group_rows:
            # 查询每个分组的分布数据
            dim_value = row.dimension or "未知"
            group_conditions = conditions + [dimension_col == row.dimension]

            # 问题分类分布
            cat_stmt = select(
                QualityTraceIndex.issue_category,
                func.count().label("cnt"),
            ).group_by(QualityTraceIndex.issue_category)
            for c in group_conditions:
                cat_stmt = cat_stmt.where(c)
            cat_result = await db.execute(cat_stmt)
            category_distribution = {r[0] or "未知": r[1] for r in cat_result.all()}

            # 处理动作分布（v1.2: 原 order_type_distribution）
            ot_stmt = select(
                QualityTraceIndex.order_type,
                func.count().label("cnt"),
            ).group_by(QualityTraceIndex.order_type)
            for c in group_conditions:
                ot_stmt = ot_stmt.where(c)
            ot_result = await db.execute(ot_stmt)
            action_type_distribution = {r[0] or "未知": r[1] for r in ot_result.all()}

            # 紧急度分布
            urg_stmt = select(
                QualityTraceIndex.urgency_level,
                func.count().label("cnt"),
            ).group_by(QualityTraceIndex.urgency_level)
            for c in group_conditions:
                urg_stmt = urg_stmt.where(c)
            urg_result = await db.execute(urg_stmt)
            urgency_distribution = {r[0] or "未知": r[1] for r in urg_result.all()}

            # 高紧急数量
            high_count_stmt = select(func.count()).where(
                QualityTraceIndex.urgency_level == "High_Priority"
            )
            for c in group_conditions:
                high_count_stmt = high_count_stmt.where(c)
            high_result = await db.execute(high_count_stmt)
            high_priority_count = high_result.scalar() or 0

            groups.append({
                "dimension": dim_value,
                "ticket_count": row.ticket_count or 0,
                "archived_count": row.archived_count or 0,
                "category_distribution": category_distribution,
                "action_type_distribution": action_type_distribution,
                "urgency_distribution": urgency_distribution,
                "avg_resolution_days": round(row.avg_resolution_days, 1) if row.avg_resolution_days else None,
                "high_priority_count": high_priority_count,
            })

        return {
            "summary": summary,
            "groups": groups,
        }

    async def get_dashboard_data(self, db: AsyncSession, query: QualityDashboardQuery) -> dict:
        """质量看板数据"""
        conditions = []
        if query.start_date:
            start_dt = datetime.strptime(query.start_date, "%Y-%m-%d")
            conditions.append(QualityTraceIndex.created_at >= start_dt)
        if query.end_date:
            end_dt = datetime.strptime(query.end_date, "%Y-%m-%d") + timedelta(days=1)
            conditions.append(QualityTraceIndex.created_at < end_dt)
        if query.model_number:
            conditions.append(QualityTraceIndex.model_number == query.model_number)

        # 1. 概览数据
        overview_stmt = select(
            func.count(QualityTraceIndex.ticket_id.distinct()).label("total_tickets"),
            func.count(QualityTraceIndex.ticket_id).label("total_archived"),
            func.count(case((QualityTraceIndex.urgency_level == "High_Priority", 1))).label("high_urgency_count"),
            func.count(case((QualityTraceIndex.archive_complete == 1, 1))).label("archive_complete_count"),
        )
        if conditions:
            overview_stmt = overview_stmt.where(*conditions)
        overview_result = await db.execute(overview_stmt)
        overview_row = overview_result.one()

        total_tickets = overview_row.total_tickets or 0
        total_archived = overview_row.total_archived or 0
        high_urgency_count = overview_row.high_urgency_count or 0
        archive_complete_count = overview_row.archive_complete_count or 0

        # v1.2: 归档完整率（替代原 SLA 达标率）
        archive_complete_rate = round(archive_complete_count / total_archived * 100, 1) if total_archived > 0 else 0

        overview = {
            "total_tickets": total_tickets,
            "total_archived": total_archived,
            "high_urgency_rate": round(high_urgency_count / total_tickets * 100, 1) if total_tickets > 0 else 0,
            "archive_complete_rate": archive_complete_rate,
        }

        # 2. 日趋势
        trend_stmt = select(
            func.date_format(QualityTraceIndex.created_at, "%Y-%m-%d").label("date"),
            func.count(QualityTraceIndex.ticket_id.distinct()).label("ticket_count"),
            func.count(QualityTraceIndex.ticket_id).label("archived_count"),
        ).group_by(func.date_format(QualityTraceIndex.created_at, "%Y-%m-%d")).order_by(func.date_format(QualityTraceIndex.created_at, "%Y-%m-%d"))
        if conditions:
            trend_stmt = trend_stmt.where(*conditions)
        trend_result = await db.execute(trend_stmt)
        trend_rows = trend_result.all()

        trend = {
            "dates": [r.date for r in trend_rows],
            "ticket_counts": [r.ticket_count for r in trend_rows],
            "archived_counts": [r.archived_count for r in trend_rows],
        }

        # 3. TOP5型号
        top_models_stmt = select(
            QualityTraceIndex.model_number,
            func.count().label("count"),
        ).group_by(QualityTraceIndex.model_number).order_by(func.count().desc()).limit(5)
        if conditions:
            top_models_stmt = top_models_stmt.where(*conditions)
        top_models_result = await db.execute(top_models_stmt)
        top_models = [
            {"model_number": r.model_number or "未知", "count": r.count}
            for r in top_models_result.all()
        ]

        # 4. 问题分类分布
        cat_stmt = select(
            QualityTraceIndex.issue_category,
            func.count().label("count"),
        ).group_by(QualityTraceIndex.issue_category)
        if conditions:
            cat_stmt = cat_stmt.where(*conditions)
        cat_result = await db.execute(cat_stmt)
        category_distribution = {r.issue_category or "未知": r.count for r in cat_result.all()}

        # 5. 处理动作分布（v1.2: 原 order_type_distribution）
        ot_stmt = select(
            QualityTraceIndex.order_type,
            func.count().label("count"),
        ).group_by(QualityTraceIndex.order_type)
        if conditions:
            ot_stmt = ot_stmt.where(*conditions)
        ot_result = await db.execute(ot_stmt)
        action_type_distribution = {r.order_type or "未知": r.count for r in ot_result.all()}

        return {
            "overview": overview,
            "trend": trend,
            "top_models": top_models,
            "category_distribution": category_distribution,
            "action_type_distribution": action_type_distribution,
        }

    async def export_report(self, db: AsyncSession, query: QualityExportQuery) -> bytes:
        """导出质量报表（返回文件字节）"""
        conditions = []
        if query.model_number:
            conditions.append(QualityTraceIndex.model_number == query.model_number)
        if query.batch_code:
            conditions.append(QualityTraceIndex.batch_code == query.batch_code)
        if query.start_date:
            start_dt = datetime.strptime(query.start_date, "%Y-%m-%d")
            conditions.append(QualityTraceIndex.created_at >= start_dt)
        if query.end_date:
            end_dt = datetime.strptime(query.end_date, "%Y-%m-%d") + timedelta(days=1)
            conditions.append(QualityTraceIndex.created_at < end_dt)

        # 查询明细数据
        stmt = select(QualityTraceIndex).order_by(QualityTraceIndex.created_at.desc())
        if conditions:
            stmt = stmt.where(*conditions)
        result = await db.execute(stmt)
        records = list(result.scalars().all())

        # 表头
        headers = [
            "工单编号", "产品型号", "批次号",
            "问题分类", "紧急度", "处理动作", "质保状态",
            "解决天数", "归档完整性", "追溯日期",
        ]

        # 中文映射表
        _CATEGORY_MAP = {
            "Missing_Parts": "配件缺失", "Operation_Error": "操作错误",
            "Software_Bug": "软件缺陷", "Hardware_Malfunction": "硬件故障",
            "Hardware_Thermal_Runaway": "热失控", "Electrical_Leakage": "漏电问题",
            "Batch_Defect": "批次缺陷", "Safety_Hazard": "安全隐患", "Other": "其他",
        }
        _URGENCY_MAP = {
            "High_Priority": "高紧急", "Medium_Priority": "中紧急", "Low_Priority": "低紧急",
        }
        _ACTION_MAP = {
            "Auto_Reply": "自动回复", "Routed": "已路由",
            "Manual_Resolved": "人工解决", "Escalated": "已升级",
        }
        _WARRANTY_MAP = {
            "In_Warranty": "保内", "Out_of_Warranty": "保外", "Unknown": "未知",
        }

        # 数据行
        rows = []
        for r in records:
            rows.append([
                r.ticket_id or "",
                r.model_number or "",
                r.batch_code or "",
                _CATEGORY_MAP.get(r.issue_category, r.issue_category or ""),
                _URGENCY_MAP.get(r.urgency_level, r.urgency_level or ""),
                _ACTION_MAP.get(r.order_type, r.order_type or ""),
                _WARRANTY_MAP.get(r.warranty_status, r.warranty_status or ""),
                str(r.resolution_days) if r.resolution_days is not None else "",
                "完整" if r.archive_complete == 1 else "部分缺失",
                r.trace_date.strftime("%Y-%m-%d") if r.trace_date else "",
            ])

        # 根据格式生成文件
        if query.format == "xlsx":
            try:
                return self._export_xlsx(headers, rows)
            except ImportError:
                # openpyxl 未安装，自动降级为 CSV
                return self._export_csv(headers, rows)
        else:
            return self._export_csv(headers, rows)

    def _export_xlsx(self, headers: list, rows: list) -> bytes:
        """使用openpyxl导出Excel"""
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "质量追溯报表"

        # 写入表头
        ws.append(headers)

        # 设置表头样式
        from openpyxl.styles import Font, Alignment, PatternFill
        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=11, color="FFFFFF")
        for cell in ws[1]:
            cell.font = header_font_white
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # 写入数据
        for row in rows:
            ws.append(row)

        # 自动调整列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except Exception:
                    pass
            adjusted_width = min(max_length + 4, 50)
            ws.column_dimensions[column].width = adjusted_width

        # 保存到字节流
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

    def _export_csv(self, headers: list, rows: list) -> bytes:
        """使用csv模块导出CSV（UTF-8 BOM编码）"""
        buffer = io.StringIO()
        # 写入UTF-8 BOM
        buffer.write("\ufeff")

        writer = csv.writer(buffer)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(row)

        return buffer.getvalue().encode("utf-8")
