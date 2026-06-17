"""质量追溯服务 QualityService 单元测试"""
import csv
import io
import re
from unittest.mock import AsyncMock, MagicMock, patch

import pytest

from backend.app.services.quality_service import QualityService
from backend.app.schemas.quality import (
    QualityTraceQuery, QualityDashboardQuery, QualityExportQuery,
)


@pytest.fixture
def service():
    return QualityService()


# ============================================================
# 纯单元测试 - 导出方法
# ============================================================

class TestExportCsv:
    """CSV导出测试"""

    def test_csv_has_bom_header(self, service):
        """CSV应包含UTF-8 BOM"""
        headers = ["工单编号", "型号"]
        rows = [["CS-001", "Pro-Max-V2"]]
        result = service._export_csv(headers, rows)
        assert result[:3] == b"\xef\xbb\xbf", "CSV应以UTF-8 BOM开头"

    def test_csv_content_correct(self, service):
        """CSV内容正确"""
        headers = ["工单编号", "型号", "批次"]
        rows = [["CS-001", "Pro-Max-V2", "B2026-03"]]
        result = service._export_csv(headers, rows)
        text = result.decode("utf-8")
        # 去掉BOM
        text = text.lstrip("\ufeff")
        lines = text.strip().split("\n")
        assert len(lines) == 2  # 表头+1行数据
        assert "工单编号" in lines[0]
        assert "CS-001" in lines[1]
        assert "Pro-Max-V2" in lines[1]

    def test_csv_empty_rows(self, service):
        """空数据行时只有表头"""
        headers = ["工单编号", "型号"]
        rows = []
        result = service._export_csv(headers, rows)
        text = result.decode("utf-8").lstrip("\ufeff")
        lines = text.strip().split("\n")
        assert len(lines) == 1  # 只有表头

    def test_csv_multiple_rows(self, service):
        """多行数据"""
        headers = ["工单编号"]
        rows = [["CS-001"], ["CS-002"], ["CS-003"]]
        result = service._export_csv(headers, rows)
        text = result.decode("utf-8").lstrip("\ufeff")
        lines = text.strip().split("\n")
        assert len(lines) == 4  # 表头+3行


class TestExportXlsx:
    """Excel导出测试"""

    def test_xlsx_returns_bytes(self, service):
        """Excel导出返回字节"""
        headers = ["工单编号", "型号"]
        rows = [["CS-001", "Pro-Max-V2"]]
        result = service._export_xlsx(headers, rows)
        assert isinstance(result, bytes)
        assert len(result) > 0

    def test_xlsx_valid_file(self, service):
        """导出的Excel是有效文件"""
        from openpyxl import load_workbook

        headers = ["工单编号", "型号", "批次"]
        rows = [["CS-001", "Pro-Max-V2", "B2026-03"]]
        result = service._export_xlsx(headers, rows)

        wb = load_workbook(io.BytesIO(result))
        ws = wb.active
        assert ws.title == "质量追溯报表"
        # 表头
        assert ws.cell(1, 1).value == "工单编号"
        assert ws.cell(1, 2).value == "型号"
        # 数据行
        assert ws.cell(2, 1).value == "CS-001"
        assert ws.cell(2, 2).value == "Pro-Max-V2"

    def test_xlsx_empty_rows(self, service):
        """空数据行时Excel只有表头"""
        from openpyxl import load_workbook

        headers = ["工单编号"]
        rows = []
        result = service._export_xlsx(headers, rows)
        wb = load_workbook(io.BytesIO(result))
        ws = wb.active
        assert ws.cell(1, 1).value == "工单编号"
        assert ws.cell(2, 1).value is None  # 无数据行


class TestExportReportFormat:
    """报表导出格式选择测试"""

    @pytest.mark.asyncio
    async def test_export_csv_format(self, service):
        """选择csv格式时调用_export_csv"""
        mock_db = MagicMock()
        mock_result = MagicMock()
        mock_result.scalars.return_value.all.return_value = []
        mock_db.execute = AsyncMock(return_value=mock_result)

        query = QualityExportQuery(format="csv")
        with patch.object(service, "_export_csv", return_value=b"csv_data") as mock_csv:
            with patch.object(service, "_export_xlsx", return_value=b"xlsx_data") as mock_xlsx:
                result = await service.export_report(mock_db, query)
                assert result == b"csv_data"
                mock_csv.assert_called_once()
                mock_xlsx.assert_not_called()

    @pytest.mark.asyncio
    async def test_export_xlsx_format(self, service):
        """选择xlsx格式时调用_export_xlsx"""
        mock_db = MagicMock()
        mock_result = MagicMock()
        mock_result.scalars.return_value.all.return_value = []
        mock_db.execute = AsyncMock(return_value=mock_result)

        query = QualityExportQuery(format="xlsx")
        with patch.object(service, "_export_csv", return_value=b"csv_data") as mock_csv:
            with patch.object(service, "_export_xlsx", return_value=b"xlsx_data") as mock_xlsx:
                result = await service.export_report(mock_db, query)
                assert result == b"xlsx_data"
                mock_xlsx.assert_called_once()
                mock_csv.assert_not_called()


# ============================================================
# 数据库集成测试
# ============================================================

@pytest.mark.integration
class TestQualityServiceIntegration:
    """质量追溯服务数据库集成测试"""

    @pytest.fixture(autouse=True)
    def check_db(self, check_db_ready):
        if not check_db_ready:
            pytest.skip("数据库不可用")

    @pytest.mark.asyncio
    async def test_get_trace_data_default(self, service, db_session):
        """测试质量追溯默认查询"""
        query = QualityTraceQuery(group_by="model")
        result = await service.get_trace_data(db_session, query)
        assert "summary" in result
        assert "groups" in result
        assert "total_records" in result["summary"]
        assert "total_tickets" in result["summary"]
        assert "total_orders" in result["summary"]
        assert isinstance(result["groups"], list)

    @pytest.mark.asyncio
    async def test_get_trace_data_filter_by_model(self, service, db_session):
        """测试按型号筛选追溯数据"""
        query = QualityTraceQuery(model_number="Pro-Max-V2", group_by="model")
        result = await service.get_trace_data(db_session, query)
        for group in result["groups"]:
            assert group["dimension"] in ("Pro-Max-V2", "未知", None)

    @pytest.mark.asyncio
    async def test_get_trace_data_group_by_category(self, service, db_session):
        """测试按问题分类分组"""
        query = QualityTraceQuery(group_by="category")
        result = await service.get_trace_data(db_session, query)
        for group in result["groups"]:
            assert "dimension" in group
            assert "ticket_count" in group
            assert "order_count" in group
            assert "category_distribution" in group
            assert "order_type_distribution" in group
            assert "urgency_distribution" in group
            assert "high_priority_count" in group

    @pytest.mark.asyncio
    async def test_get_dashboard_data(self, service, db_session):
        """测试质量看板数据"""
        query = QualityDashboardQuery()
        result = await service.get_dashboard_data(db_session, query)
        assert "overview" in result
        assert "trend" in result
        assert "top_models" in result
        assert "category_distribution" in result
        assert "order_type_distribution" in result

        overview = result["overview"]
        assert "total_tickets" in overview
        assert "total_orders" in overview
        assert "high_urgency_rate" in overview
        assert "sla_rate" in overview

        trend = result["trend"]
        assert "dates" in trend
        assert "ticket_counts" in trend
        assert "order_counts" in trend

    @pytest.mark.asyncio
    async def test_export_report_csv(self, service, db_session):
        """测试CSV报表导出"""
        query = QualityExportQuery(format="csv")
        result = await service.export_report(db_session, query)
        assert isinstance(result, bytes)
        assert result[:3] == b"\xef\xbb\xbf"  # BOM

    @pytest.mark.asyncio
    async def test_export_report_xlsx(self, service, db_session):
        """测试Excel报表导出"""
        query = QualityExportQuery(format="xlsx")
        result = await service.export_report(db_session, query)
        assert isinstance(result, bytes)
        assert len(result) > 0
        assert result[:2] == b"PK"

    @pytest.mark.asyncio
    async def test_dashboard_high_urgency_rate_calculation(self, service, db_session):
        """测试高紧急率计算"""
        query = QualityDashboardQuery()
        result = await service.get_dashboard_data(db_session, query)
        overview = result["overview"]
        total = overview["total_tickets"]
        rate = overview["high_urgency_rate"]
        if total > 0:
            assert 0 <= rate <= 100, f"高紧急率应在0-100之间, 实际: {rate}"
        else:
            assert rate == 0
