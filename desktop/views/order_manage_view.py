import os
import tempfile

from PyQt6.QtWidgets import (
    QWidget, QVBoxLayout, QHBoxLayout, QLabel, QPushButton,
    QComboBox, QTableWidget, QTableWidgetItem, QHeaderView,
    QTabBar, QSizePolicy, QDialog, QFormLayout, QTextEdit,
    QDialogButtonBox, QMessageBox, QMenu, QAbstractItemView
)
from PyQt6.QtCore import Qt, QThread, pyqtSignal
from PyQt6.QtGui import QFont

from desktop.api_client import ApiClient
from shared.constants import ORDER_STATUS_TRANSITIONS, OrderStatus, OrderType

# 单据类型标签映射
ORDER_TYPE_LABELS = {
    "Replacement": "补发单",
    "Repair": "维修单",
    "Return_Exchange": "退换单",
    "Tech_Support": "技术支援单",
    "QC": "质检单",
}

# 单据状态标签映射
ORDER_STATUS_LABELS = {
    "pending": "待处理",
    "processing": "处理中",
    "executing": "执行中",
    "completed": "已完成",
    "cancelled": "已取消",
}

# 单据状态颜色映射 (背景色, 前景色)
ORDER_STATUS_COLORS = {
    "pending": ("#FFF3CD", "#856404"),
    "processing": ("#D6EAF8", "#2471A3"),
    "executing": ("#E8DAEF", "#7D3C98"),
    "completed": ("#D5F5E3", "#1E8449"),
    "cancelled": ("#F2F3F4", "#BDC3C7"),
}

# 状态Tab列表
STATUS_TABS = [
    ("全部", ""),
    ("待处理", "pending"),
    ("处理中", "processing"),
    ("执行中", "executing"),
    ("已完成", "completed"),
    ("已取消", "cancelled"),
]


class OrderStatusBadge(QLabel):
    """单据状态标签"""

    def __init__(self, status, parent=None):
        super().__init__(parent)
        bg, fg = ORDER_STATUS_COLORS.get(status, ("#F2F3F4", "#7F8C8D"))
        label = ORDER_STATUS_LABELS.get(status, status or "未知")
        self.setText(f"  {label}  ")
        self.setStyleSheet(
            f"background-color: {bg}; color: {fg}; "
            f"border-radius: 10px; padding: 4px 12px; font-weight: bold; font-size: 12px;"
        )
        self.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.setFixedHeight(26)


class LoadOrdersThread(QThread):
    """加载单据列表线程"""
    finished = pyqtSignal(object)

    def __init__(self, api_client, order_type=None, status=None, page=1, page_size=20):
        super().__init__()
        self.api_client = api_client
        self.order_type = order_type
        self.status = status
        self.page = page
        self.page_size = page_size

    def run(self):
        result = self.api_client.get_orders(
            order_type=self.order_type,
            status=self.status,
            page=self.page,
            page_size=self.page_size,
        )
        self.finished.emit(result)


class ActionThread(QThread):
    """通用操作线程"""
    finished = pyqtSignal(object)

    def __init__(self, func, *args, **kwargs):
        super().__init__()
        self.func = func
        self.args = args
        self.kwargs = kwargs

    def run(self):
        try:
            result = self.func(*self.args, **self.kwargs)
            self.finished.emit(result)
        except Exception:
            self.finished.emit(None)


class OrderDetailDialog(QDialog):
    """单据详情弹窗"""

    def __init__(self, order_data, parent=None):
        super().__init__(parent)
        self.setWindowTitle("单据详情")
        self.setMinimumWidth(500)
        self.setMinimumHeight(400)

        layout = QVBoxLayout(self)

        form = QFormLayout()
        form.setSpacing(10)

        # 基本信息
        form.addRow("单据编号：", QLabel(order_data.get("order_no", "-")))
        form.addRow("关联工单：", QLabel(order_data.get("ticket_id", "-")))

        order_type = order_data.get("order_type", "")
        type_label = ORDER_TYPE_LABELS.get(order_type, order_type or "-")
        form.addRow("单据类型：", QLabel(type_label))

        status = order_data.get("status", "")
        status_badge = OrderStatusBadge(status)
        form.addRow("状态：", status_badge)

        form.addRow("部门：", QLabel(order_data.get("department", "-")))
        form.addRow("SLA时限：", QLabel(f"{order_data.get('sla_hours', '-')}小时"))

        deadline = order_data.get("deadline", "-")
        if deadline and len(str(deadline)) >= 16:
            deadline = str(deadline)[:16].replace("T", " ")
        form.addRow("截止时间：", QLabel(str(deadline)))

        form.addRow("产品型号：", QLabel(order_data.get("model_number", "-")))
        form.addRow("批次号：", QLabel(order_data.get("batch_code", "-")))
        form.addRow("备注：", QLabel(order_data.get("remark", "-") or "无"))

        created = order_data.get("created_at", "-")
        if created and len(str(created)) >= 16:
            created = str(created)[:16].replace("T", " ")
        form.addRow("创建时间：", QLabel(str(created)))

        layout.addLayout(form)

        close_btn = QPushButton("关闭")
        close_btn.clicked.connect(self.accept)
        layout.addWidget(close_btn, alignment=Qt.AlignmentFlag.AlignCenter)


class OrderStatusChangeDialog(QDialog):
    """更新单据状态弹窗"""

    def __init__(self, current_status, parent=None):
        super().__init__(parent)
        self.setWindowTitle("更新单据状态")
        self.setMinimumWidth(360)
        self._result = None

        layout = QVBoxLayout(self)

        form = QFormLayout()
        self.status_combo = QComboBox()

        # 根据当前状态获取可流转的状态
        transitions = ORDER_STATUS_TRANSITIONS.get(OrderStatus(current_status), [])
        for ts in transitions:
            self.status_combo.addItem(
                ORDER_STATUS_LABELS.get(ts.value, ts.value), ts.value
            )
        if self.status_combo.count() == 0:
            self.status_combo.addItem("无可用状态", "")
        form.addRow("新状态：", self.status_combo)

        self.note_edit = QTextEdit()
        self.note_edit.setPlaceholderText("备注（选填）")
        self.note_edit.setMaximumHeight(80)
        form.addRow("备注：", self.note_edit)

        layout.addLayout(form)

        btn_box = QDialogButtonBox(
            QDialogButtonBox.StandardButton.Ok | QDialogButtonBox.StandardButton.Cancel
        )
        btn_box.accepted.connect(self._on_accept)
        btn_box.rejected.connect(self.reject)
        layout.addWidget(btn_box)

    def _on_accept(self):
        self._result = {
            "status": self.status_combo.currentData(),
            "note": self.note_edit.toPlainText().strip() or None,
        }
        self.accept()

    def get_result(self):
        return self._result


class OrderManageView(QWidget):
    """出单管理视图"""

    def __init__(self, api_client: ApiClient, role=None, parent=None):
        super().__init__(parent)
        self.api_client = api_client
        self._role = role
        self._current_page = 1
        self._total_pages = 1
        self._total_count = 0
        self._orders = []
        self._current_status = ""
        self._current_order_type = ""
        self._load_thread = None
        self._action_thread = None
        self._setup_ui()
        self._load_orders()

    def _setup_ui(self):
        main_layout = QVBoxLayout(self)
        main_layout.setContentsMargins(24, 24, 24, 24)
        main_layout.setSpacing(16)

        # 顶部标题栏
        header_layout = QHBoxLayout()
        title = QLabel("出单管理")
        title.setStyleSheet("font-size: 20px; font-weight: bold; color: #2C3E50;")
        header_layout.addWidget(title)
        header_layout.addStretch()
        main_layout.addLayout(header_layout)

        # 筛选栏：类型下拉框
        filter_layout = QHBoxLayout()
        filter_layout.setSpacing(12)

        filter_layout.addWidget(QLabel("单据类型："))
        self.type_combo = QComboBox()
        self.type_combo.addItem("全部", "")
        self.type_combo.addItem("补发单", "Replacement")
        self.type_combo.addItem("维修单", "Repair")
        self.type_combo.addItem("退换单", "Return_Exchange")
        self.type_combo.addItem("技术支援单", "Tech_Support")
        self.type_combo.addItem("质检单", "QC")
        self.type_combo.setFixedWidth(160)
        filter_layout.addWidget(self.type_combo)

        self.search_btn = QPushButton("查询")
        self.search_btn.clicked.connect(self._on_search)
        filter_layout.addWidget(self.search_btn)

        self.reset_btn = QPushButton("重置")
        self.reset_btn.clicked.connect(self._on_reset)
        filter_layout.addWidget(self.reset_btn)

        filter_layout.addStretch()

        self.refresh_btn = QPushButton("刷新")
        self.refresh_btn.clicked.connect(self._on_refresh)
        filter_layout.addWidget(self.refresh_btn)

        main_layout.addLayout(filter_layout)

        # 状态Tab栏
        self.status_tab = QTabBar()
        self.status_tab.setExpanding(False)
        for label, status_value in STATUS_TABS:
            self.status_tab.addTab(label)
        self.status_tab.currentChanged.connect(self._on_status_tab_changed)
        main_layout.addWidget(self.status_tab)

        # 单据列表表格
        self.table = QTableWidget()
        self.table.setColumnCount(6)
        self.table.setHorizontalHeaderLabels(
            ["单据编号", "类型", "关联工单", "部门", "状态", "SLA时限"]
        )
        self.table.horizontalHeader().setSectionResizeMode(0, QHeaderView.ResizeMode.Fixed)
        self.table.setColumnWidth(0, 160)
        self.table.horizontalHeader().setSectionResizeMode(1, QHeaderView.ResizeMode.Fixed)
        self.table.setColumnWidth(1, 100)
        self.table.horizontalHeader().setSectionResizeMode(2, QHeaderView.ResizeMode.Fixed)
        self.table.setColumnWidth(2, 160)
        self.table.horizontalHeader().setSectionResizeMode(3, QHeaderView.ResizeMode.Fixed)
        self.table.setColumnWidth(3, 100)
        self.table.horizontalHeader().setSectionResizeMode(4, QHeaderView.ResizeMode.Fixed)
        self.table.setColumnWidth(4, 100)
        self.table.horizontalHeader().setSectionResizeMode(5, QHeaderView.ResizeMode.Stretch)
        self.table.setAlternatingRowColors(True)
        self.table.setSelectionBehavior(QAbstractItemView.SelectionBehavior.SelectRows)
        self.table.setSelectionMode(QAbstractItemView.SelectionMode.SingleSelection)
        self.table.verticalHeader().setVisible(False)
        self.table.setEditTriggers(QAbstractItemView.EditTrigger.NoEditTriggers)
        self.table.cellDoubleClicked.connect(self._on_row_double_clicked)
        self.table.setContextMenuPolicy(Qt.ContextMenuPolicy.CustomContextMenu)
        self.table.customContextMenuRequested.connect(self._on_context_menu)
        main_layout.addWidget(self.table, 1)

        # 底部分页栏
        page_layout = QHBoxLayout()
        self.total_label = QLabel("共 0 条记录")
        self.total_label.setStyleSheet("color: #7F8C8D; font-size: 13px;")
        page_layout.addWidget(self.total_label)
        page_layout.addStretch()

        self.prev_btn = QPushButton("上一页")
        self.prev_btn.clicked.connect(self._on_prev_page)
        page_layout.addWidget(self.prev_btn)

        self.page_label = QLabel("第 1 / 1 页")
        self.page_label.setStyleSheet("color: #2C3E50; font-size: 13px; padding: 0 12px;")
        page_layout.addWidget(self.page_label)

        self.next_btn = QPushButton("下一页")
        self.next_btn.clicked.connect(self._on_next_page)
        page_layout.addWidget(self.next_btn)

        page_layout.addStretch()

        self.export_btn = QPushButton("导出Excel")
        self.export_btn.clicked.connect(self._on_export)
        page_layout.addWidget(self.export_btn)

        main_layout.addLayout(page_layout)

    def _load_orders(self):
        """加载单据列表"""
        self._set_controls_enabled(False)

        self._load_thread = LoadOrdersThread(
            self.api_client,
            order_type=self._current_order_type or None,
            status=self._current_status or None,
            page=self._current_page,
            page_size=20,
        )
        self._load_thread.finished.connect(self._on_orders_loaded)
        self._load_thread.start()

    def _set_controls_enabled(self, enabled):
        """设置控件启用/禁用状态"""
        self.search_btn.setEnabled(enabled)
        self.reset_btn.setEnabled(enabled)
        self.refresh_btn.setEnabled(enabled)
        self.export_btn.setEnabled(enabled)
        self.prev_btn.setEnabled(enabled and self._current_page > 1)
        self.next_btn.setEnabled(enabled and self._current_page < self._total_pages)

    def _on_orders_loaded(self, result):
        """单据列表加载完成"""
        self._set_controls_enabled(True)

        if result is None:
            self.table.setRowCount(0)
            self.total_label.setText("无法加载单据数据")
            self.total_label.setStyleSheet("color: #FF4444; font-size: 13px;")
            return

        total = result.get("total", 0)
        self._total_count = total
        self._total_pages = max(1, (total + 19) // 20)
        self._orders = result.get("data", [])

        self.total_label.setText(f"共 {total} 条记录")
        self.total_label.setStyleSheet("color: #7F8C8D; font-size: 13px;")
        self.page_label.setText(f"第 {self._current_page} / {self._total_pages} 页")

        self.prev_btn.setEnabled(self._current_page > 1)
        self.next_btn.setEnabled(self._current_page < self._total_pages)

        self._populate_table()

    def _populate_table(self):
        """填充表格数据"""
        self.table.setRowCount(len(self._orders))
        for row, order in enumerate(self._orders):
            # 单据编号
            item_no = QTableWidgetItem(order.get("order_no", "-"))
            self.table.setItem(row, 0, item_no)

            # 类型
            order_type = order.get("order_type", "")
            type_text = ORDER_TYPE_LABELS.get(order_type, order_type or "-")
            item_type = QTableWidgetItem(type_text)
            self.table.setItem(row, 1, item_type)

            # 关联工单
            item_ticket = QTableWidgetItem(order.get("ticket_id", "-"))
            self.table.setItem(row, 2, item_ticket)

            # 部门
            item_dept = QTableWidgetItem(order.get("department", "-"))
            self.table.setItem(row, 3, item_dept)

            # 状态
            status = order.get("status", "")
            status_badge = OrderStatusBadge(status)
            self.table.setCellWidget(row, 4, status_badge)

            # SLA时限
            sla_hours = order.get("sla_hours", "-")
            sla_text = f"{sla_hours}h" if sla_hours else "-"
            item_sla = QTableWidgetItem(sla_text)
            self.table.setItem(row, 5, item_sla)

        for row in range(self.table.rowCount()):
            self.table.setRowHeight(row, 36)

    def _on_status_tab_changed(self, index):
        """状态Tab切换"""
        if 0 <= index < len(STATUS_TABS):
            self._current_status = STATUS_TABS[index][1]
            self._current_page = 1
            self._load_orders()

    def _on_search(self):
        """查询按钮点击"""
        self._current_order_type = self.type_combo.currentData()
        self._current_page = 1
        self._load_orders()

    def _on_reset(self):
        """重置筛选条件"""
        self.type_combo.setCurrentIndex(0)
        self.status_tab.setCurrentIndex(0)
        self._current_order_type = ""
        self._current_status = ""
        self._current_page = 1
        self._load_orders()

    def _on_refresh(self):
        """刷新数据"""
        self._load_orders()

    def _on_prev_page(self):
        """上一页"""
        if self._current_page > 1:
            self._current_page -= 1
            self._load_orders()

    def _on_next_page(self):
        """下一页"""
        if self._current_page < self._total_pages:
            self._current_page += 1
            self._load_orders()

    def _on_row_double_clicked(self, row, col):
        """双击行查看详情"""
        if row < 0 or row >= len(self._orders):
            return
        order = self._orders[row]
        dialog = OrderDetailDialog(order, self)
        dialog.exec()

    def _on_context_menu(self, pos):
        """右键菜单"""
        row = self.table.rowAt(pos.y())
        if row < 0 or row >= len(self._orders):
            return

        order = self._orders[row]
        current_status = order.get("status", "")
        order_id = order.get("id")

        menu = QMenu(self)

        # 查看详情
        detail_action = menu.addAction("查看详情")
        detail_action.triggered.connect(lambda: self._on_row_double_clicked(row, 0))

        # 更新状态（根据当前状态允许的流转）
        transitions = ORDER_STATUS_TRANSITIONS.get(OrderStatus(current_status), [])
        if transitions:
            menu.addSeparator()
            update_menu = menu.addMenu("更新状态")
            for ts in transitions:
                label = ORDER_STATUS_LABELS.get(ts.value, ts.value)
                action = update_menu.addAction(label)
                action.triggered.connect(
                    lambda checked, s=ts.value, oid=order_id: self._update_order_status_direct(oid, s)
                )

        menu.exec(self.table.viewport().mapToGlobal(pos))

    def _update_order_status_direct(self, order_id, new_status):
        """直接更新单据状态（右键菜单触发）"""
        if not order_id:
            QMessageBox.warning(self, "操作失败", "无法获取单据ID")
            return

        self._set_controls_enabled(False)
        self._action_thread = ActionThread(
            self.api_client.update_order_status, order_id, new_status
        )
        self._action_thread.finished.connect(self._on_action_done)
        self._action_thread.start()

    def _on_action_done(self, result):
        """操作完成回调"""
        self._set_controls_enabled(True)
        if result is None:
            QMessageBox.warning(self, "操作失败", "操作失败，请检查网络连接或后端服务")
        elif isinstance(result, dict) and result.get("code") not in (0, None):
            QMessageBox.warning(self, "操作失败", result.get("message", "未知错误"))
        else:
            QMessageBox.information(self, "操作成功", "操作已成功执行")
            self._load_orders()

    def _on_export(self):
        """导出Excel"""
        self._set_controls_enabled(False)
        self._action_thread = ActionThread(
            self.api_client.get_orders,
            order_type=self._current_order_type or None,
            status=self._current_status or None,
            page=1,
            page_size=1000,
        )
        self._action_thread.finished.connect(self._on_export_data_loaded)
        self._action_thread.start()

    def _on_export_data_loaded(self, result):
        """导出数据加载完成，执行导出"""
        self._set_controls_enabled(True)

        if result is None:
            QMessageBox.warning(self, "导出失败", "无法获取单据数据")
            return

        orders = result.get("data", [])
        if not orders:
            QMessageBox.information(self, "提示", "没有可导出的数据")
            return

        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment, PatternFill

            wb = Workbook()
            ws = wb.active
            ws.title = "出单数据"

            headers = ["单据编号", "类型", "关联工单", "部门", "状态", "SLA时限", "产品型号", "批次号", "创建时间"]
            ws.append(headers)

            header_font = Font(bold=True, size=11, color="FFFFFF")
            header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center", vertical="center")

            for order in orders:
                order_type = order.get("order_type", "")
                type_label = ORDER_TYPE_LABELS.get(order_type, order_type or "-")
                status = order.get("status", "")
                status_label = ORDER_STATUS_LABELS.get(status, status or "-")
                created = order.get("created_at", "-")
                if created and len(str(created)) >= 16:
                    created = str(created)[:16].replace("T", " ")

                ws.append([
                    order.get("order_no", ""),
                    type_label,
                    order.get("ticket_id", ""),
                    order.get("department", ""),
                    status_label,
                    order.get("sla_hours", ""),
                    order.get("model_number", ""),
                    order.get("batch_code", ""),
                    str(created),
                ])

            for col in ws.columns:
                max_length = 0
                column = col[0].column_letter
                for cell in col:
                    try:
                        if len(str(cell.value)) > max_length:
                            max_length = len(str(cell.value))
                    except Exception:
                        pass
                adjusted_width = min(max_length + 4, 50)
                ws.column_dimensions[column].width = adjusted_width

            tmp_dir = tempfile.gettempdir()
            filepath = os.path.join(tmp_dir, "orders_export.xlsx")
            wb.save(filepath)

            QMessageBox.information(self, "导出成功", f"数据已导出到：\n{filepath}")
        except ImportError:
            QMessageBox.warning(self, "导出失败", "缺少 openpyxl 库，无法导出Excel文件")
        except Exception as e:
            QMessageBox.warning(self, "导出失败", f"导出失败：{str(e)}")
